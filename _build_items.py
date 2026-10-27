# -*- coding: utf-8 -*-
"""从《rf5简陋物品数据8_27.xlsx》生成 js/items.js（RF5 物品数据）"""
import openpyxl, json, io, os, collections

SRC = r'G:/10257/Downloads/符文工房攻略/rf5简陋物品数据8_27.xlsx'
OUT = r'G:/wit_tools/work/RF_WIKI/js/items.js'

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
base = [r for r in wb['物品基础信息'].iter_rows(values_only=True)][2:]
eq = [r for r in wb['装备信息'].iter_rows(values_only=True)][1:]
wb.close()

num = lambda v: 0 if v is None or v in ('', '-') else v
bmap = {}
for r in base:
    if r[2]:
        bmap.setdefault(r[2], r)

# ---------- 武器 ----------
WEP = ['单手剑', '双手剑', '枪', '双剑', '拳套', '斧', '锤', '魔杖']
weapons = []
for r in eq:
    if r[5] in WEP and r[2]:
        weapons.append([r[2], r[5], r[6] or '无',
                        num(r[7]), num(r[9]), num(r[11]), num(r[14]), num(r[16])])

# ---------- 防具 / 饰品 ----------
RESN = [('火抗', 26), ('水抗', 27), ('地抗', 28), ('风抗', 29), ('光抗', 30),
        ('暗抗', 31), ('爱抗', 32), ('无抗', 33),
        ('毒抗', 37), ('封抗', 38), ('麻抗', 39), ('眠抗', 40),
        ('疲抗', 41), ('病抗', 42), ('昏厥抗', 43), ('吸收抗', 44)]
SHORT = {'火抗': '火', '水抗': '水', '地抗': '地', '风抗': '风', '光抗': '光', '暗抗': '暗',
         '爱抗': '爱', '无抗': '无', '毒抗': '毒', '封抗': '封', '麻抗': '麻', '眠抗': '眠',
         '疲抗': '疲', '病抗': '病', '昏厥抗': '昏厥', '吸收抗': '吸收'}


def ressum(r):
    out = []
    for name, idx in RESN:
        v = num(r[idx])
        if isinstance(v, (int, float)) and v:
            out.append(f'{SHORT[name]}{v:g}')
    return '、'.join(out)


ARM = ['盾', '帽子', '铠甲', '鞋子']
armors = []
for r in eq:
    if r[5] in ARM and r[2]:
        armors.append([r[2], r[5], num(r[8]), num(r[10]), num(r[34]), num(r[36]),
                       ressum(r), r[48] or ''])

BONUS = [('攻', 7), ('防', 8), ('魔攻', 9), ('魔防', 10), ('力量', 11),
         ('智力', 12), ('体力', 13)]


def bonussum(r):
    out = []
    for name, idx in BONUS:
        v = num(r[idx])
        if isinstance(v, (int, float)) and v:
            out.append(f'{name}+{v:g}')
    return ' '.join(out)


accs = []
for r in eq:
    if r[5] == '饰品' and r[2]:
        accs.append([r[2], bonussum(r), ressum(r), r[48] or ''])

# ---------- 农具 ----------
TOOL = ['锄头', '洒水壶', '镰刀', '斧(农具)', '锤(农具)', '钓竿', '毛刷', '剪刀', '放大镜']
tools = []
for r in eq:
    if r[5] in TOOL and r[2]:
        tools.append([r[2], r[5].replace('(农具)', ''), r[6] or '无',
                      num(r[7]), num(r[9]), r[48] or ''])

# ---------- 素材 ----------
MAT = ['素材', '素材(矿石)', '素材(石头)', '素材(爪牙)', '素材(碎片)', '素材(毛皮)',
       '素材(鳞片)', '素材(粉末)', '素材(结晶)', '素材(宝石)', '素材(棒茎)',
       '素材(布皮)', '素材(羽毛)', '素材(线绳)', '素材(液体)', '素材(骨甲)']
mats = []
for r in base:
    if r[5] in MAT and r[2]:
        sub = r[5].replace('素材', '').strip('()') or '其他'
        mats.append([r[2], sub, num(r[6]), num(r[7]), num(r[8])])

# ---------- 料理 ----------
DISH = {'平底锅料理': '平底锅', '刀工料理': '刀工', '火锅料理': '火锅',
        '火锅料理(果酱)': '果酱', '烤箱料理': '烤箱', '搅拌机料理': '搅拌机',
        '徒手料理': '徒手', '蒸笼料理': '蒸笼', '配方面包': '面包'}
dishes = []
for r in base:
    if r[5] in DISH and r[2]:
        dishes.append([r[2], DISH[r[5]], num(r[6]), num(r[7]), num(r[8])])

# ---------- 食材 ----------
ING = {'蔬菜': '蔬菜', '蔬菜(大)': '蔬菜', '蔬菜(蘑菇)': '蔬菜',
       '水果': '水果', '水果(大)': '水果', '水果(流星)': '水果',
       '花草': '花草', '花草(大)': '花草',
       '乳制品': '乳制品', '乳制品(牛奶)': '乳制品', '乳制品(鸡蛋)': '乳制品',
       '料理材料': '料理材料'}
ings = []
for r in base:
    if r[5] in ING and r[2]:
        ings.append([r[2], ING[r[5]], num(r[6]), num(r[7]), num(r[8])])

# ---------- 药品 ----------
pots = []
for r in base:
    if r[5] in ('药', '营养剂') and r[2]:
        pots.append([r[2], r[5], num(r[6]), num(r[7]), num(r[8])])


def dump(name, rows, comment):
    return (f'/* {comment} */\nconst {name} = ' +
            json.dumps(rows, ensure_ascii=False, separators=(',', ':')) + ';\n')


out = io.StringIO()
out.write('/* ============================================================\n')
out.write('   RF5 物品数据（武器 / 防具 / 农具 / 素材 / 料理 / 药品）\n')
out.write('   来源：《rf5简陋物品数据8_27.xlsx》数据解包表，汉化版字段。\n')
out.write('   ============================================================ */\n\n')
out.write(dump('EQW_RF5', weapons, '武器：名称, 种类, 属性, 物攻, 魔攻, 力量, 暴击, 晕眩'))
out.write('\n')
out.write(dump('EQA_RF5', armors, '防具：名称, 部位, 物防, 魔防, 暴击抗, 晕眩抗, 抗性摘要, 特效'))
out.write('\n')
out.write(dump('EQC_RF5', accs, '饰品：名称, 数值摘要, 抗性摘要, 特效'))
out.write('\n')
out.write(dump('EQT_RF5', tools, '农具：名称, 类型, 属性, 物攻, 魔攻, 特效'))
out.write('\n')
out.write(dump('MAT_RF5', mats, '素材：名称, 分类, 买价, 卖价, 稀有度'))
out.write('\n')
out.write(dump('DSH_RF5', dishes, '料理：名称, 做法, 买价, 卖价, 稀有度'))
out.write('\n')
out.write(dump('ING_RF5', ings, '食材：名称, 类别, 买价, 卖价, 稀有度'))
out.write('\n')
out.write(dump('POT_RF5', pots, '药品：名称, 分类, 买价, 卖价, 稀有度'))

with open(OUT, 'w', encoding='utf-8') as f:
    f.write(out.getvalue())

print('weapons', len(weapons))
print('armors', len(armors), collections.Counter(a[1] for a in armors))
print('accs', len(accs))
print('tools', len(tools), collections.Counter(t[1] for t in tools))
print('mats', len(mats), collections.Counter(m[1] for m in mats))
print('dishes', len(dishes), collections.Counter(d[1] for d in dishes))
print('ings', len(ings), collections.Counter(i[1] for i in ings))
print('pots', len(pots))
print('bytes', os.path.getsize(OUT))
